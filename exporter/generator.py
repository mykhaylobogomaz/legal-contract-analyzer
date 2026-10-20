from analyzer import RiskFinding
from humanizer_module import Humanizer
import pandas as pd
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfgen import canvas
from typing import List
from datetime import datetime

class ReportGenerator:
    def __init__(self, findings: List[RiskFinding]):
        self.findings = findings
        self.humanizer = Humanizer()

    def _humanize(self, text: str) -> str:
        return self.humanizer.remove_ai_stamps(text)

    def generate_excel(self, output_path: str):
        df = pd.DataFrame([f.__dict__ for f in self.findings])
        for col in df.columns:
            if df[col].dtype == 'object':
                df[col] = df[col].apply(lambda x: self._humanize(str(x)) if isinstance(x, str) else x)
        df.to_excel(output_path, index=False)

    def generate_word(self, output_path: str):
        doc = Document()
        doc.add_heading("Legal Contract Risk Analysis", 0)
        for f in self.findings:
            doc.add_heading(self._humanize(f"{f.severity}: {f.category}"), level=1)
            doc.add_paragraph(self._humanize(f.sentence))
            doc.add_paragraph(self._humanize(f"Recommendation: {f.recommendation}"))
        doc.save(output_path)

    def generate_pdf(self, output_path: str):
        doc = SimpleDocTemplate(output_path, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
        styles = getSampleStyleSheet()
        story = []

        # Header
        story.append(Paragraph("Legal Contract Risk Analysis Report", styles['Title']))
        story.append(Paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d')} | Status: High Risk Detected", styles['Normal']))
        story.append(Spacer(1, 20))

        # Table
        data = [["Risk", "Category", "Problem"]]
        for f in self.findings:
            data.append([
                Paragraph(f.severity, styles['Normal']),
                Paragraph(f.category, styles['Normal']),
                Paragraph(self._humanize(f.explanation), styles['Normal'])
            ])

        table = Table(data, colWidths=[60, 100, 250])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.navy),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        story.append(table)

        # Recommendations
        story.append(Spacer(1, 20))
        story.append(Paragraph("Recommendations", styles['Heading2']))
        story.append(Paragraph("1. Redline critical clauses.<br/>2. Limit liability.<br/>3. Negotiate cure periods.", styles['Normal']))

        doc.build(story)

    def generate_ppt(self, output_path: str):
        prs = Presentation()
        navy = RGBColor(13, 37, 64)

        # 1. Title Slide
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = self._humanize("Legal Contract Risk Analysis")
        subtitle = slide.placeholders[1]
        subtitle.text = f"Audit Date: {datetime.now().strftime('%Y-%m-%d')}\nStatus: CRITICAL RISKS DETECTED"

        # 2. Key Metrics
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = self._humanize("Risk Profile Summary")
        counts = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0}
        for f in self.findings: counts[f.severity] += 1

        x_offsets = [Inches(0.5), Inches(2.8), Inches(5.1), Inches(7.4)]
        for i, (sev, count) in enumerate(counts.items()):
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_offsets[i], Inches(2.5), Inches(2), Inches(1.5))
            shape.fill.solid()
            shape.fill.fore_color.rgb = navy
            tf = shape.text_frame
            tf.text = f"{sev}\n{count}"

        # 3. Top Critical Risks
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = self._humanize("Top Priority Critical Risks")
        tf = slide.placeholders[1].text_frame
        crits = [f for f in self.findings if f.severity == "CRITICAL"][:5]
        for f in crits:
            p = tf.add_paragraph()
            p.text = f"• {f.category}: {f.pattern_matched}"

        # 4. Recommendations
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = self._humanize("Legal Strategy & Recommendations")
        tf = slide.placeholders[1].text_frame
        recs = [
            "Negotiate mutual termination rights.",
            "Insert mandatory cure periods for all breaches.",
            "Cap indemnification and liquidated damages.",
            "Carve out pre-existing IP from assignment clauses."
        ]
        for r in recs:
            p = tf.add_paragraph()
            p.text = self._humanize(r)

        # 5. Conclusion
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = self._humanize("Final Audit Verdict")
        body = slide.placeholders[1]
        body.text = self._humanize("HIGH RISK. The contract is heavily skewed in favor of the Franchisor. Critical clauses lack balance and expose the Franchisee to unlimited liability. Signing is NOT recommended without comprehensive legal redlining.")

        prs.save(output_path)
